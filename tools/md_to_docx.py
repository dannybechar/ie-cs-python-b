# Convert a simple Markdown file (headings, paragraphs, bullets, tables, code blocks) to .docx.
# Hebrew paragraphs and tables become right-to-left; code blocks stay left-to-right.
# Usage: python md_to_docx.py input.md output.docx [extra.py appended as a code block] [--worksheet]
# --worksheet: code blocks become boxed grey squares, all code is bold Courier New, left-aligned and left-to-right;
#              Hebrew paragraphs and Hebrew table cells are right-to-left and right-aligned.
import re
import sys

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_LINE_SPACING
from docx.shared import Cm, Inches, Pt

HEB = re.compile("[\u0590-\u05ff]")
CODE_FONT = "Courier New"  # repository rule: all code is Courier New, left-to-right
WORKSHEET = False  # set by --worksheet: boxed bold code


def set_font(run, name):
    run.font.name = name
    rfonts = run._r.get_or_add_rPr().get_or_add_rFonts()
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(attr), name)


def set_rtl(par):
    ppr = par._p.get_or_add_pPr()
    ppr.append(OxmlElement("w:bidi"))
    if not WORKSHEET:
        par.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    # worksheet mode: no explicit alignment. A right-to-left paragraph starts at the right edge in Word and
    # LibreOffice alike; an explicit "right" is read as the opposite edge by Word in bidi paragraphs.


def add_runs(par, text, rtl, font=None):
    for i, chunk in enumerate(re.split(r"(\*\*.+?\*\*|`[^`]+`)", text)):
        if not chunk:
            continue
        bold = chunk.startswith("**")
        code = chunk.startswith("`")
        chunk = chunk.strip("*") if bold else chunk.strip("`") if code else chunk
        run = par.add_run(chunk)
        run.bold = bold or (code and WORKSHEET)
        if code or font:
            set_font(run, font or CODE_FONT)
        if rtl and not code:
            rpr = run._r.get_or_add_rPr()
            rpr.append(OxmlElement("w:rtl"))


def paragraph(doc, text, style=None):
    text = text.replace("&nbsp;", " ")
    rtl = bool(HEB.search(text))
    par = doc.add_paragraph(style=style)
    add_runs(par, text, rtl)
    if rtl:
        set_rtl(par)
    return par


def boxed_code_block(doc, lines):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl.autofit = False
    tblpr = tbl._tbl.tblPr
    width = OxmlElement("w:tblW")
    width.set(qn("w:w"), "5000")
    width.set(qn("w:type"), "pct")
    tblpr.append(width)
    trpr = tbl.rows[0]._tr.get_or_add_trPr()
    trpr.append(OxmlElement("w:cantSplit"))  # keep the whole box on one page
    cell = tbl.cell(0, 0)
    tcpr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    for side in ("top", "left", "bottom", "right"):
        b = OxmlElement(f"w:{side}")
        b.set(qn("w:val"), "single")
        b.set(qn("w:sz"), "12")
        b.set(qn("w:color"), "000000")
        borders.append(b)
    tcpr.append(borders)
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), "F2F2F2")
    tcpr.append(shd)
    first = True
    for line in lines:
        par = cell.paragraphs[0] if first else cell.add_paragraph()
        first = False
        par.paragraph_format.space_after = Pt(0)
        par.paragraph_format.space_before = Pt(0)
        par.alignment = WD_ALIGN_PARAGRAPH.LEFT
        run = par.add_run(line if line else " ")
        set_font(run, CODE_FONT)
        run.bold = True
        run.font.size = Pt(11)
        par.paragraph_format.keep_with_next = True  # a box stays with what follows it (an answer box, a table)
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(0)
    spacer.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    spacer.paragraph_format.line_spacing = Pt(8)
    spacer.paragraph_format.keep_with_next = True


def code_block(doc, lines):
    if WORKSHEET:
        boxed_code_block(doc, lines)
        return
    for line in lines:
        par = doc.add_paragraph()
        par.paragraph_format.space_after = Pt(0)
        par.alignment = WD_ALIGN_PARAGRAPH.LEFT
        run = par.add_run(line if line else " ")
        set_font(run, CODE_FONT)
        run.font.size = Pt(10)
    doc.add_paragraph()


def table(doc, rows):
    cells = [[c.strip() for c in r.strip().strip("|").split("|")] for r in rows]
    cells = [r for r in cells if not all(re.fullmatch(r":?-+:?", c) for c in r)]
    tbl = doc.add_table(rows=len(cells), cols=len(cells[0]), style="Table Grid")
    rtl = any(HEB.search(c) for r in cells for c in r)
    if rtl:
        tbl._tbl.tblPr.append(OxmlElement("w:bidiVisual"))
    for i, r in enumerate(cells):
        for j, c in enumerate(r):
            cell = tbl.cell(i, j)
            cell.text = ""
            par = cell.paragraphs[0]
            for k, part in enumerate(c.split("<br>")):  # <br> = line break inside a cell
                if k:
                    par.add_run().add_break()
                add_runs(par, part.replace("&nbsp;", " ").replace("\\_", "_"), bool(HEB.search(c)))
            if not HEB.search(c):
                par.alignment = WD_ALIGN_PARAGRAPH.LEFT  # code / English cells: left-to-right
            elif WORKSHEET:
                set_rtl(par)  # Hebrew cells: right-to-left, right-aligned
            if i == 0:
                for run in par.runs:
                    run.bold = True
            if WORKSHEET and re.fullmatch(r"(&nbsp;|<br>)+", c.replace(" ", "")):
                par.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
                par.paragraph_format.line_spacing = Pt(20)  # room for handwriting in answer cells
            if WORKSHEET and re.fullmatch(r"\d+", c):
                par.alignment = WD_ALIGN_PARAGRAPH.CENTER  # numbers (round numbers, expected values)
    if WORKSHEET:
        for row in tbl.rows:
            row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))  # a row never splits across pages
        weights = []
        for j in range(len(cells[0])):
            longest = max(len(re.sub(r"<br>|&nbsp;|`", "", r[j])) for r in cells)
            weights.append(min(max(longest, 14), 40))
        total = sum(weights)
        tbl.autofit = False
        for j, w in enumerate(weights):
            tbl.columns[j].width = Inches(6.69 * w / total)  # the table grid (LibreOffice uses this)
            for row in tbl.rows:
                row.cells[j].width = Inches(6.69 * w / total)  # and every cell (Word uses this)
        for row in tbl.rows[:-1]:
            for cell in row.cells:
                for cp in cell.paragraphs:
                    cp.paragraph_format.keep_with_next = True  # keep the whole table on one page
    doc.add_paragraph()


def convert(src, dst, extra=None):
    doc = Document()
    doc.styles["Normal"].font.name = "Arial"
    doc.styles["Normal"].font.size = Pt(11)
    if WORKSHEET:  # A4, 2 cm margins -> 17 cm (6.7 in) of usable width
        sec = doc.sections[0]
        sec.page_width, sec.page_height = Cm(21), Cm(29.7)
        sec.left_margin = sec.right_margin = sec.top_margin = sec.bottom_margin = Cm(2)
    lines = open(src, encoding="utf-8").read().splitlines()
    i = 0
    pending_break = False
    prev = None  # the last paragraph written; kept on the same page as a code box that follows it
    while i < len(lines):
        line = lines[i]
        if line.startswith("```"):
            if WORKSHEET and prev is not None:
                prev.paragraph_format.keep_with_next = True
            prev = None
            j = i + 1
            while not lines[j].startswith("```"):
                j += 1
            code_block(doc, lines[i + 1:j])
            i = j + 1
            continue
        if line.startswith("|"):
            if WORKSHEET and prev is not None:
                prev.paragraph_format.keep_with_next = True
            prev = None
            j = i
            while j < len(lines) and lines[j].startswith("|"):
                j += 1
            table(doc, lines[i:j])
            i = j
            continue
        m = re.match(r"(#+) (.*)", line)
        if m:
            h = doc.add_heading(level=min(len(m.group(1)), 3))
            if pending_break:
                h.paragraph_format.page_break_before = True
                pending_break = False
            add_runs(h, m.group(2), bool(HEB.search(m.group(2))))
            if HEB.search(m.group(2)):
                set_rtl(h)
        elif re.match(r"\s*- ", line):
            paragraph(doc, re.sub(r"^\s*- ", "", line), "List Bullet")
            prev = None
        elif re.match(r"\d+\. ", line):
            prev = paragraph(doc, line)
        elif line.strip() == "<pagebreak>":
            pending_break = True  # the next heading starts a new page (invisible in Markdown renderers)
            prev = None
        elif line.strip() in ("", "---") or line.startswith("<"):
            pass
        else:
            prev = paragraph(doc, line)
        i += 1
    if extra:
        doc.add_heading("Model solutions (Python)", level=2)
        code_block(doc, open(extra, encoding="utf-8").read().splitlines())
    doc.save(dst)


if __name__ == "__main__":
    args = [a for a in sys.argv[1:] if a != "--worksheet"]
    WORKSHEET = "--worksheet" in sys.argv
    convert(args[0], args[1], args[2] if len(args) > 2 else None)
